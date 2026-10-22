#!/usr/bin/env python3
"""Regression fixtures for observable merge/content failures; no user files."""
from __future__ import annotations
import copy
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import parse_xml
from lxml import etree as E
from PIL import Image

from merge_integrity import Package, W, M, R, q, digest, scaffold, verify, IntegrityError
from merge_pipeline import build, release
from apply_review_comments import apply
from audit_review_comments import audit


def make_doc(path, texts):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    doc.save(path)
    return doc


def patch(path, mutate):
    with ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    root = E.fromstring(files["word/document.xml"])
    mutate(root)
    files["word/document.xml"] = E.tostring(root)
    with ZipFile(path, "w", ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)


class MergeRegression(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="merge-regression-", dir=Path.cwd())
        self.root = Path(self.temp.name)
        self.template = self.root / "template.docx"
        self.source = self.root / "source.docx"
        self.output = self.root / "output.docx"
        self.plan_path = self.root / "plan.json"
        make_doc(self.template, ["模板正文原型", "插入位置"])
        make_doc(self.source, ["第一项研究内容。", "第二项研究内容。"])
        shutil.copyfile(self.source, self.output)

    def tearDown(self):
        self.temp.cleanup()

    def plan(self, sources=None):
        sources = sources or [self.source]
        plan = scaffold(self.template, sources)
        plan["canonical_sections"] = ["主要内容", "总体方案"]
        plan["merged_docx"] = str(self.output)
        target = list(Package(self.output).inventory())
        for index, entry in enumerate(plan["content_register"]):
            entry["section"] = "主要内容"
            entry["targets"] = [target[index % len(target)]]
        return plan

    def write_plan(self, plan):
        self.plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        return self.plan_path

    def report(self, plan):
        return verify(self.write_plan(plan))

    def test_baseline_roundtrip(self):
        self.assertTrue(self.report(self.plan())["passed"])

    def test_missing_version_fails(self):
        plan = self.plan()
        del plan["plan_version"]
        with self.assertRaises(IntegrityError):
            self.report(plan)

    def test_legacy_version_cannot_reenable_deletion(self):
        plan = self.plan()
        plan["plan_version"] = 1
        with self.assertRaises(IntegrityError):
            self.report(plan)

    def test_missing_source_unit_fails(self):
        plan = self.plan()
        plan["content_register"].pop()
        report = self.report(plan)
        self.assertFalse(report["passed"])
        self.assertTrue(any("unregistered source block" in e for e in report["errors"]))

    def test_empty_register_fails(self):
        plan = self.plan()
        plan["content_register"] = []
        with self.assertRaises(IntegrityError):
            self.report(plan)

    def test_repeated_registration_fails(self):
        plan = self.plan()
        plan["content_register"].append(copy.deepcopy(plan["content_register"][0]))
        self.assertFalse(self.report(plan)["passed"])

    def test_changed_frozen_source_fails(self):
        plan = self.plan()
        make_doc(self.source, ["源文件已经改变"])
        with self.assertRaises(IntegrityError):
            self.report(plan)

    def test_formula_structure_is_not_flat_text(self):
        doc = make_doc(self.source, ["公式"])
        doc.paragraphs[0]._p.append(parse_xml(f'<m:oMath xmlns:m="{M}"><m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f></m:oMath>'))
        doc.save(self.source)
        shutil.copyfile(self.source, self.output)
        plan = self.plan()
        def mutate(root):
            old = root.find('.//{' + M + '}oMath')
            old.getparent().replace(old, E.fromstring(f'<m:oMath xmlns:m="{M}"><m:sSup><m:e><m:r><m:t>a</m:t></m:r></m:e><m:sup><m:r><m:t>b</m:t></m:r></m:sup></m:sSup></m:oMath>'))
        patch(self.output, mutate)
        self.assertFalse(self.report(plan)["passed"])

    def test_swapped_pictures_fail_at_actual_location(self):
        doc = Document()
        for i, color in enumerate(["red", "blue"]):
            image = self.root / f"image-{i}.png"
            Image.new("RGB", (12, 12), color).save(image)
            p = doc.add_paragraph(f"图片段落 {i}")
            p.add_run().add_picture(str(image))
        doc.save(self.source)
        shutil.copyfile(self.source, self.output)
        plan = self.plan()
        def mutate(root):
            nodes = root.xpath('//*[@r:embed]', namespaces={"r": R})
            left, right = [n.get("{" + R + "}embed") for n in nodes]
            nodes[0].set("{" + R + "}embed", right)
            nodes[1].set("{" + R + "}embed", left)
        patch(self.output, mutate)
        self.assertFalse(self.report(plan)["passed"])

    def test_superscript_loss_fails(self):
        doc = Document()
        run = doc.add_paragraph().add_run("2")
        run.font.superscript = True
        doc.save(self.source)
        shutil.copyfile(self.source, self.output)
        plan = self.plan()
        def mutate(root):
            for n in root.iter(q("vertAlign")):
                n.getparent().remove(n)
        patch(self.output, mutate)
        self.assertFalse(self.report(plan)["passed"])

    def test_picture_proportional_resize_allowed_but_stretch_rejected(self):
        image = self.root / "picture.png"
        Image.new("RGB", (24, 12), "green").save(image)
        doc = Document()
        doc.add_picture(str(image))
        doc.save(self.source)
        shutil.copyfile(self.source, self.output)
        plan = self.plan()
        def shrink(root):
            for node in root.xpath('//*[@cx and @cy]'):
                node.set("cx", str(int(node.get("cx")) // 2))
                node.set("cy", str(int(node.get("cy")) // 2))
        patch(self.output, shrink)
        self.assertTrue(self.report(plan)["passed"])
        def stretch(root):
            for node in root.xpath('//*[@cx and @cy]'):
                node.set("cx", str(int(node.get("cx")) * 2))
        patch(self.output, stretch)
        self.assertFalse(self.report(plan)["passed"])

    def test_fabricated_equal_hashes_do_not_prove_duplicate(self):
        plan = self.plan()
        item = plan["content_register"][1]
        item.update(retention="exact_duplicate", covered_by=plan["content_register"][0]["id"], targets=[], source_fingerprint="same", canonical_fingerprint="same")
        self.assertFalse(self.report(plan)["passed"])

    def test_actual_duplicate_is_accepted(self):
        doc = Document()
        doc.add_paragraph("完全相同的内容")
        doc.add_paragraph("完全相同的内容")
        doc.save(self.source)
        make_doc(self.output, ["完全相同的内容"])
        plan = self.plan()
        plan["content_register"][1].update(retention="exact_duplicate", covered_by=plan["content_register"][0]["id"], targets=[])
        self.assertTrue(self.report(plan)["passed"])

    def test_shadow_needs_user_authorization(self):
        plan = self.plan()
        plan["content_register"][1].update(retention="replace_shadow", covered_by=plan["content_register"][0]["id"], targets=[])
        self.assertFalse(self.report(plan)["passed"])
        plan["content_register"][1]["user_authorization"] = {"instruction": "删除第二段简写，以第一段为准", "decision_ref": "fixture-user-request"}
        self.assertFalse(self.report(plan)["passed"])
        make_doc(self.output, ["第一项研究内容。"])
        self.assertTrue(self.report(plan)["passed"])

    def test_synthesis_does_not_remove_preservation_duty(self):
        plan = self.plan()
        entry = plan["content_register"][0]
        entry["uses"] = [{"kind": "synthesis", "section": "总体方案", "targets": entry["targets"], "evidence": "源内容作为总体论述依据"}]
        self.assertTrue(self.report(plan)["passed"])
        entry["targets"] = []
        self.assertFalse(self.report(plan)["passed"])

    def test_old_synthesis_status_rejected(self):
        plan = self.plan()
        plan["content_register"][0]["retention"] = "synthesis_evidence"
        self.assertFalse(self.report(plan)["passed"])

    def test_split_preserves_content_and_requires_comment(self):
        make_doc(self.source, ["标题：正文内容。"])
        make_doc(self.output, ["标题：", "正文内容。"])
        plan = self.plan()
        targets = list(Package(self.output).inventory())
        item = plan["content_register"][0]
        item.update(targets=targets, structural_group="split-1")
        plan["structural_changes"] = [{"id": "split-1", "operation": "split", "entries": [item["id"]], "targets": targets, "authority": "源稿冒号与语法边界", "comment_id": "decision-1"}]
        self.assertFalse(self.report(plan)["passed"])
        plan["review_comments"]["comments"] = [{"id": "decision-1", "anchor_contains": "标题：", "comment_text": "[结构调整] 将源段按明确边界分为标题和正文，请确认。"}]
        self.assertTrue(self.report(plan)["passed"])

    def test_table_cell_swap_fails_even_with_remapped_text(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        doc.save(self.source)
        shutil.copyfile(self.source, self.output)
        plan = self.plan()
        def mutate(root):
            nodes = list(root.iter(q("t")))
            nodes[0].text, nodes[1].text = nodes[1].text, nodes[0].text
        patch(self.output, mutate)
        entries = [e for e in plan["content_register"] if e["source_preview"] in {"A", "B"}]
        entries[0]["targets"], entries[1]["targets"] = entries[1]["targets"], entries[0]["targets"]
        self.assertFalse(self.report(plan)["passed"])

    def test_citation_without_definition_fails(self):
        make_doc(self.source, ["正文[1]"])
        make_doc(self.output, ["正文[9]"])
        plan = self.plan()
        plan["numbering_map"] = [{"id": "c1", "kind": "citation", "research_doc_id": "research-01", "old": "[1]", "new": "[9]"}]
        plan["content_register"][0]["text_changes"] = [{"mapping_id": "c1"}]
        self.assertFalse(self.report(plan)["passed"])

    def test_scoped_comment_on_repeated_title(self):
        make_doc(self.source, ["重复标题", "重复标题"])
        units = list(Package(self.source).inventory())
        manifest = {"mode": "review", "comments": [{"id": "d1", "anchor_contains": "重复标题", "anchor_block_id": units[1], "category": "待裁定", "comment_text": "[待裁定] 请确认第二处标题。"}]}
        apply(self.source, manifest, self.output)
        path = self.root / "comments.json"
        path.write_text(json.dumps(manifest), encoding="utf-8")
        self.assertTrue(audit(self.output, path)["passed"])
        comments = Package(self.output).root("word/document.xml").findall('.//' + q("p"))
        self.assertIsNone(comments[0].find(q("commentRangeStart")))
        self.assertIsNotNone(comments[1].find(q("commentRangeStart")))

    def test_ambiguous_comment_fails_before_writing(self):
        make_doc(self.source, ["重复标题", "重复标题"])
        manifest = {"mode": "review", "comments": [{"id": "d1", "anchor_contains": "重复标题", "comment_text": "请确认"}]}
        with self.assertRaises(IntegrityError):
            apply(self.source, manifest, self.root / "new-output.docx")
        self.assertFalse((self.root / "new-output.docx").exists())

    def test_pipeline_build_and_stale_release_gate(self):
        plan = self.plan()
        adapter = self.root / "adapter.py"
        adapter.write_text('''from merge_integrity import Package
from template_writer import TemplateWriter

def build(plan, candidate_path):
    writer = TemplateWriter(plan["template"]["file"])
    keys = list(writer.anchors)
    refs = []
    for descriptor in plan["research_docs"]:
        source = Package(descriptor["file"])
        for entry in plan["content_register"]:
            if entry["research_doc_id"] == descriptor["id"]:
                node = writer.write_paragraph(source, entry["source_block"], anchor=keys[-1], prototype=keys[0])
                refs.append((entry, node))
    for entry, node in refs:
        entry["targets"] = [writer.target_id(node)]
    writer.save(candidate_path)
    return plan
''', encoding="utf-8")
        result = build(self.write_plan(plan), adapter, self.root / "work")
        self.assertEqual(result["status"], "candidate_ready_for_review")
        gates = self.root / "work/review-gates.json"
        with self.assertRaises(IntegrityError):
            release(Path(result["plan"]), Path(result["candidate"]), gates, self.root / "released.docx", "review")
        gate_value = json.loads(gates.read_text())
        gate_value["candidate_sha256"] = "outdated"
        gates.write_text(json.dumps(gate_value))
        with self.assertRaisesRegex(IntegrityError, "earlier candidate"):
            release(Path(result["plan"]), Path(result["candidate"]), gates, self.root / "released.docx", "review")
        self.assertFalse((self.root / "released.docx").exists())

    def test_release_requires_current_plan_and_real_evidence_files(self):
        plan = self.plan()
        self.write_plan(plan)
        note = self.root / "fixture-review.txt"
        note.write_text("Synthetic evidence for release-protocol testing, not a visual assessment.", encoding="utf-8")
        png = self.root / "fixture-page.png"
        Image.new("RGB", (16, 16), "white").save(png)
        evidence = {"file": str(note), "sha256": digest(note.read_bytes())}
        from merge_pipeline import GATES
        gates = {name: {"status": "passed", "reviewed_by": "protocol-fixture", "finding": "Synthetic protocol assertion", "evidence": [evidence]} for name in GATES}
        gates["visual_pages"].update(page_count=1, pages=[{"page": 1, "file": str(png), "sha256": digest(png.read_bytes()), "status": "passed"}])
        reviews = {"candidate_sha256": digest(self.output.read_bytes()), "plan_sha256": digest(self.plan_path.read_bytes()), "gates": gates}
        gates_path = self.root / "gates.json"
        gates_path.write_text(json.dumps(reviews), encoding="utf-8")
        result = release(self.plan_path, self.output, gates_path, self.root / "release-ok.docx", "review")
        self.assertEqual(result["sha256"], digest(self.output.read_bytes()))
        plan["decisions"] = []
        self.write_plan(plan)
        with self.assertRaisesRegex(IntegrityError, "earlier merge plan"):
            release(self.plan_path, self.output, gates_path, self.root / "release-stale.docx", "review")
        reviews["plan_sha256"] = digest(self.plan_path.read_bytes())
        gates_path.write_text(json.dumps(reviews), encoding="utf-8")
        note.write_text("Changed evidence", encoding="utf-8")
        with self.assertRaisesRegex(IntegrityError, "stale review evidence"):
            release(self.plan_path, self.output, gates_path, self.root / "release-stale-evidence.docx", "review")


if __name__ == "__main__":
    unittest.main(verbosity=2)
